"""Операции над файлом: проверить, исправить, добавить примечания. Работают с байтами DOCX."""

from __future__ import annotations

import logging
from collections import OrderedDict
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from . import fixes as F
from .analyzer import Analyzer
from .model import Analysis

log = logging.getLogger(__name__)

_TEXT_CHANGING = {
    "heading_caps", "heading_number_dot", "heading_end_dot", "heading_style", "toc_title",
    "fig_caption_text", "fig_number", "tbl_caption_text", "tbl_number", "lst_caption_text", "lst_number",
}


class BadDocument(Exception):
    pass


def _load(data: bytes):
    try:
        return Document(BytesIO(data))
    except Exception as exc:  # повреждённый архив, .doc, переименованный в .docx, и т.п.
        raise BadDocument(str(exc)) from exc


def _save(document) -> bytes:
    out = BytesIO()
    document.save(out)
    return out.getvalue()


def check(data: bytes) -> Analysis:
    return Analyzer(_load(data)).run()


@dataclass
class FixResult:
    data: bytes
    before: Analysis
    after: Analysis
    applied: int
    failed: int


def _apply(analysis: Analysis) -> tuple[int, int, bool]:
    status: dict[int, bool] = {}
    applied = failed = 0
    changed_text = False
    for issue in analysis.issues:
        if issue.fix is None:
            continue
        key = id(issue.fix)
        if key not in status:
            try:
                issue.fix()
                status[key] = True
            except Exception:
                log.exception("Не удалось применить исправление %s", issue.code)
                status[key] = False
        if status[key]:
            applied += 1
            changed_text = changed_text or issue.code in _TEXT_CHANGING
        else:
            failed += 1
    for post in analysis.post_fixes:
        try:
            post()
        except Exception:
            log.exception("Ошибка при обновлении ссылок")
    return applied, failed, changed_text


def fix(data: bytes, max_passes: int = 3) -> FixResult:
    """Применяет автоисправления. Несколько проходов: одно исправление (например, смена стиля
    заголовка) может открыть мелкие несоответствия, которые снимает следующий проход."""
    document = _load(data)
    before = None
    applied = failed = 0
    changed_text = has_toc = False
    for pass_no in range(max_passes):
        analyzer = Analyzer(document)
        analysis = analyzer.run()
        if before is None:
            before = analysis
        has_toc = has_toc or analyzer.has_toc_field
        n_applied, n_failed, changed = _apply(analysis)
        if pass_no == 0:
            applied, failed = n_applied, n_failed
        changed_text = changed_text or changed
        F.normalize_styles(document)
        if n_applied == 0:
            break
    if has_toc and changed_text:
        F.request_fields_update(document)
    result = _save(document)
    after = check(result)
    return FixResult(result, before, after, applied, failed)


def annotate(data: bytes) -> tuple[bytes, Analysis]:
    """Копия документа с примечаниями Word на абзацах, где найдены ошибки."""
    document = _load(data)
    analysis = Analyzer(document).run()
    body = document.element.body
    groups: OrderedDict[int, tuple[object, list]] = OrderedDict()
    for issue in analysis.issues:
        anchor = issue.anchor
        if anchor is None or not _inside(anchor, body):
            continue
        groups.setdefault(id(anchor), (anchor, []))[1].append(issue)

    for anchor, issues in groups.values():
        paragraph = Paragraph(anchor, document._body)
        runs = [Run(r, paragraph) for r in anchor.findall(qn("w:r"))]
        if not runs:
            runs = [paragraph.add_run()]
        lines, seen = [], set()
        for issue in issues:
            mark = "✖" if issue.is_error else "⚠"
            line = f"{mark} {issue.rule.title}"
            if issue.detail:
                line += f": {issue.detail}"
            if line in seen:
                continue
            seen.add(line)
            lines.append(line)
            lines.append(f"   → {issue.rule.hint}")
        try:
            document.add_comment(runs=[runs[0], runs[-1]], text="\n".join(lines[:30]),
                                 author="DoxiBot", initials="DB")
        except Exception:
            log.exception("Не удалось добавить примечание")
    return _save(document), analysis


def _inside(el, root) -> bool:
    while el is not None:
        if el is root:
            return True
        el = el.getparent()
    return False
