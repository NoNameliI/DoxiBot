"""Генерирует тестовые отчёты: bad.docx (много нарушений) и good.docx (оформлен по правилам)."""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Mm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from doxibot import fixes as F  # noqa: E402

OUT = Path(__file__).resolve().parent / "samples"

OMML = (
    '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
    '<m:oMath><m:r><m:t>E=m</m:t></m:r><m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e>'
    '<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>'
)


def png_bytes() -> BytesIO:
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (400, 250), (120, 120, 140)).save(buf, "PNG")
    buf.seek(0)
    return buf


def base_normal(doc, good: bool):
    st = doc.styles["Normal"]
    if good:
        st.font.name = "Times New Roman"
        st.font.size = Pt(14)
        pf = st.paragraph_format
        pf.first_line_indent = Cm(1.25)
        pf.line_spacing = 1.5
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for lvl in (1, 2):
            h = doc.styles[f"Heading {lvl}"]
            h.font.name = "Times New Roman"
            h.font.size = Pt(14)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            rf = h.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            h.paragraph_format.space_before = Pt(0)
            h.paragraph_format.space_after = Pt(0)
            h.paragraph_format.line_spacing = 1.5


def heading(doc, text, level=1, center=False, good=True):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if good:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(0) if center else Cm(1.25)
    return p


def picture(doc, good=True):
    p = doc.add_paragraph()
    p.add_run().add_picture(png_bytes(), width=Cm(8))
    if good:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def caption(doc, text, center=True):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    return p


def table(doc, header=True, size=None):
    t = doc.add_table(rows=3, cols=3)
    t.style = doc.styles["Table Grid"]
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.text = (f"Столбец {j + 1}" if header else str(j + 1)) if i == 0 else str(i * 10 + j)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.size = Pt(size or 12)
    return t


def page_numbers(doc):
    F.add_page_numbers(doc)


def build_good() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)
        s.header_distance = s.footer_distance = Cm(1.25)
    base_normal(doc, True)
    t = doc.add_paragraph("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Отчёт по практике").runs[0].font.size = Pt(20)
    doc.add_page_break()

    heading(doc, "СОДЕРЖАНИЕ", center=True)
    toc = doc.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    for kind, payload in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '), ("separate", None),
                          ("text", "ВВЕДЕНИЕ\t3"), ("end", None)):
        toc._p.append(F._field_run(kind, payload))

    heading(doc, "ВВЕДЕНИЕ", center=True)
    doc.add_paragraph("В работе рассмотрены виды и типы котов, приведённые на рисунке 1.1 и в таблице 1.1.")
    heading(doc, "1 ВИДЫ")
    doc.add_paragraph("Внешний вид кота показан в соответствии с рисунком 1.1.")
    picture(doc)
    caption(doc, "Рисунок 1.1 – Серый кот")
    doc.add_paragraph("Характеристики приведены в таблице 1.1.")
    caption(doc, "Таблица 1.1 – Пример", center=False)
    table(doc)
    doc.add_paragraph("Энергия вычисляется по формуле (1.1).")
    doc.add_paragraph()
    f = doc.add_paragraph()
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.paragraph_format.first_line_indent = Cm(0)
    f._p.append(parse_xml(OMML))
    n = doc.add_paragraph("(1.1)")
    n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    n.paragraph_format.first_line_indent = Cm(0)
    heading(doc, "1.1 Тип первый", level=2)
    doc.add_paragraph("Для достижения цели необходимо выполнить следующие задачи:")
    heading(doc, "ЗАКЛЮЧЕНИЕ", center=True)
    doc.add_paragraph("Цель работы достигнута.")
    heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", center=True)
    doc.add_paragraph("1. Рейтинг регионов [электронный ресурс]. – URL: http://example.ru (дата обращения 24.09.2018).")
    page_numbers(doc)
    return doc


def build_bad() -> Document:
    doc = Document()  # поля по умолчанию 1 дюйм
    base_normal(doc, False)
    doc.add_paragraph("ТИТУЛЬНЫЙ ЛИСТ")
    doc.add_page_break()

    heading(doc, "Содержание", good=False)  # не капсом, не по центру, при главах «ГЛАВА» — должно быть ОГЛАВЛЕНИЕ
    heading(doc, "ВВЕДЕНИЕ.", good=False)
    p = doc.add_paragraph("Основной текст в Calibri, без отступа и с интервалом после. См. рис. 1.2.")
    p.runs[0].bold = True
    u = p.add_run(" Подчёркнутая синяя ссылка")
    u.underline = True
    u.font.color.rgb = RGBColor(0, 0, 255)

    heading(doc, "Глава 1. Виды", good=False)
    doc.add_paragraph("Текст главы со ссылкой на рисунок 1.1.1.")
    picture(doc, good=False)
    caption(doc, "Рис. 1.1. Серый кот.", center=False)
    picture(doc, good=True)
    caption(doc, "Рисунок 1.1.1 - Кот")
    caption(doc, "Таблица 1 Пример", center=True)
    table(doc, header=False, size=12)
    doc.add_paragraph("Ещё таблица без подписи:")
    table(doc, size=14)
    b = doc.add_paragraph("Первый пункт", style="List Bullet")
    doc.add_paragraph("Второй пункт", style="List Bullet")
    doc.add_paragraph("Текст перед формулой без пустой строки.")
    f = doc.add_paragraph()
    f._p.append(parse_xml(OMML))
    doc.add_paragraph("(5)")
    heading(doc, "1.1 ТИП ПЕРВЫЙ", level=2, good=False)
    doc.add_paragraph("Текст подглавы.")
    fl = doc.add_paragraph("Непронумерованный жирный заголовок:")
    heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", good=False)
    doc.add_paragraph("1. Сайт. – URL: http://example.ru")
    return doc


def main():
    OUT.mkdir(exist_ok=True)
    build_good().save(OUT / "good.docx")
    build_bad().save(OUT / "bad.docx")
    print("saved to", OUT)


if __name__ == "__main__":
    main()
