"""Тесты движка проверки: python -m pytest tests"""

from __future__ import annotations

import sys
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from doxibot import report, service  # noqa: E402
from doxibot.analyzer import parse_caption  # noqa: E402
from doxibot.ooxml import paragraph_text, replace_span  # noqa: E402
from tests.make_samples import build_bad, build_good  # noqa: E402


def _bytes(doc) -> bytes:
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def good() -> bytes:
    return _bytes(build_good())


@pytest.fixture(scope="module")
def bad() -> bytes:
    return _bytes(build_bad())


def codes(analysis) -> set[str]:
    return {i.code for i in analysis.issues}


def test_good_document_has_no_issues(good):
    analysis = service.check(good)
    assert analysis.issues == [], report.full_text_report(analysis, "good.docx")


def test_bad_document_detects_key_rules(bad):
    found = codes(service.check(bad))
    expected = {
        "margins", "page_numbers_missing", "font_name", "font_size", "font_color", "underline", "bold_body",
        "align_body", "indent_body", "spacing_body", "line_spacing_body", "heading_align", "heading_caps",
        "heading_not_caps", "heading_end_dot", "toc_title", "toc_manual", "picture_align", "fig_caption_text",
        "fig_number", "fig_ref_missing", "fig_ref_abbrev", "tbl_caption_text", "tbl_number", "table_no_caption",
        "table_header", "table_uniform", "formula_number", "formula_blank", "formula_number_align", "list_tab",
        "list_indent", "list_marker", "source_access_date",
    }
    assert expected <= found, expected - found


def test_fix_leaves_only_manual_issues(bad):
    result = service.fix(bad)
    assert result.failed == 0
    assert result.applied > 50
    assert all(not i.fixable for i in result.after.issues), [i.code for i in result.after.issues if i.fixable]
    fixed = Document(BytesIO(result.data))
    texts = [p.text for p in fixed.paragraphs]
    assert "Рисунок 1.1 – Серый кот" in texts
    assert "Рисунок 1.2 – Кот" in texts
    assert "Таблица 1.1 – Пример" in texts
    assert "ОГЛАВЛЕНИЕ" in texts
    assert "Текст главы со ссылкой на рисунок 1.2." in texts  # ссылка обновлена вслед за номером


def test_fix_is_stable(bad):
    once = service.fix(bad)
    twice = service.fix(once.data)
    assert twice.applied == 0


def test_annotate_adds_comments(bad):
    data, analysis = service.annotate(bad)
    doc = Document(BytesIO(data))
    assert len(list(doc.comments)) > 5


def test_bad_file_raises():
    with pytest.raises(service.BadDocument):
        service.check(b"not a docx")


@pytest.mark.parametrize("text,kind,num,dash,name,tail", [
    ("Рисунок 1.1 – Серый кот", "fig", "1.1", "–", "Серый кот", ""),
    ("Рис. 1.1. Серый кот.", "fig", "1.1.", None, "Серый кот", "."),
    ("Таблица А.2 - Данные", "tbl", "А.2", "-", "Данные", ""),
    ("Листинг 3.1 — Код", "lst", "3.1", "—", "Код", ""),
])
def test_parse_caption(text, kind, num, dash, name, tail):
    cap = parse_caption(text)
    assert (cap["kind"], cap["num"], cap["dash"], cap["name"], cap["tail"]) == (kind, num, dash, name, tail)


def test_parse_caption_ignores_sentences():
    assert parse_caption("Таблица показывает результаты") is None


def test_replace_span_across_runs():
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Рис. ")
    p.add_run("1.1")
    p.add_run(". Кот.")
    replace_span(p._p, 0, 4, "Рисунок")
    assert paragraph_text(p._p) == "Рисунок 1.1. Кот."
    replace_span(p._p, 11, 12, " –")
    assert paragraph_text(p._p) == "Рисунок 1.1 – Кот."


def test_chat_messages_fit_telegram_limit(bad):
    messages, _ = report.chat_messages(service.check(bad), "bad.docx")
    assert messages and all(len(m) <= 4096 for m in messages)
